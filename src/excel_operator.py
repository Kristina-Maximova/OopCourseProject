import os
import pandas as pd
from openpyxl.styles import NamedStyle, Font, Alignment, PatternFill, Border, Side
from openpyxl.workbook import Workbook

from src.abstract_operator import VacanciesOperator
from src.vacancy import Vacancy


class ExcelOperator(VacanciesOperator):
    """ Класс для сохранения вакансий в excel-файл """

    def __init__(self, source_name: str, source_url: str, vacancies_list: list, file_path: str = "") -> None:
        """ Конструктор класса для сохранения вакансий в excel-файл"""
        self.source_name = source_name
        self.source_url = source_url
        self.__file_path = file_path if file_path else self._get_file_name(source_name)
        self.vacancies_list = vacancies_list if vacancies_list else []
        super().__init__()

    def __enter__(self):
        """ Метод для получения данных о вакансиях из excel-файла"""
        # Если нет папки, создаем ее
        if not os.path.exists('..\\excel_files'):
            os.makedirs('..\\excel_files')
        if not os.path.exists(self.__file_path):
            # Сoздаем новый файл с нужными колонками
            df = pd.DataFrame(
                columns=['name', 'salary', 'currency', 'created_at', 'url', 'requirement', 'schedule', "vac_id"])
            df.to_excel(self.__file_path, index=False, sheet_name="sheet_1", freeze_panes=(1, 0))
        # считываем данные из exel-файла и заносим их в список объекта как вакансии
        vacs = pd.read_excel(self.__file_path, sheet_name='sheet_1')
        if not vacs.empty:
            for vacancy in vacs:
                self.vacancies_list.append(Vacancy(**vacancy))

    def __exit__(self, exc_type, exc_val, exc_tb):
        """ Метод для записи данных о вакансиях в excel-файл"""
        vac_list_dict = [vac.to_dict for vac in self.vacancies_list]
        # Создаем новую книгу Excel
        sheet_name = 'sheet_1'
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        if vac_list_dict:
            header = list(vac_list_dict[0].keys())
            ws.append(header)  # Записываем заголовки
            for row in vac_list_dict:
                ws.append([row[col] for col in header])
        # Настраиваем стили для красивого вида
        header_style = NamedStyle(name='header')
        header_style.font = Font(bold=True, color='FFFFFF')
        header_style.alignment = Alignment(horizontal='center', vertical='center')
        header_style.fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        border_style = Border(
            left=Side(border_style='thin', color='000000'),
            right=Side(border_style='thin', color='000000'),
            top=Side(border_style='thin', color='000000'),
            bottom=Side(border_style='thin', color='000000')
        )
        header_style.border = border_style

        cell_style = NamedStyle(name='cell')
        cell_style.alignment = Alignment(horizontal='left', vertical='center')
        cell_style.border = border_style
        for cell in ws[1]:  # Применяем стиль к заголовкам
            cell.style = header_style

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.style = cell_style

        # Автоматическое изменение ширины столбцов
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except Exception:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
        # Сохраняем файл
        wb.save(self.__file_path)

    def add_vacancies(self, vacs_list):
        with self:  # тут работают __enter__ и __exit__
            if vacs_list:
                for vac in vacs_list:
                    if isinstance(vac, Vacancy) and vac not in set(self.vacancies_list):
                        self.vacancies_list.append(vac)

    def get_vacancies(self, *args, **kwargs):

        pass

    def del_vacancies(self, *args, **kwargs):

        pass

    @staticmethod
    def _get_file_name(source_name: str):
        """ Метод для генерации имени файла, возвращает путь к файлу)"""
        file_name = f"{source_name.replace(' ', '_')}_vacancies.xlsx"
        project_root = os.path.dirname(os.path.abspath(__file__))
        path_to_file = os.path.join(project_root, "..", "excel_files", file_name)
        return path_to_file


if __name__ == "__main__":  # pragma: no cover
    vac1 = Vacancy(**{'name': 'Middle Backend Developer (Python FastAPI + SQL)',
                      'salary': 210000,
                      'currency': "RUR",
                      'created_at': '2025-02-21T14:46:49+0300',
                      'url': 'https://api.hh.ru/vacancies/117570423?host=hh.ru',
                      'requirement': 'Написание качественного кода и тестов к нему. '
                                     'Необходимый стек: <highlighttext>Python</highlighttext> (asyncio), '
                                     'чистый SQL (не ORM), FastAPI (весь функционал), микросервисная архитектура...',
                      'schedule': 'Удаленная работа',
                      "vac_id": 1
                      })
    vac2 = Vacancy(**{'name': 'Python-разработчик',
                      'salary': 80000,
                      'currency': "RUR",
                      'created_at': '2025-02-21T14:00:01+0300',
                      'url': 'https://api.hh.ru/vacancies/116455408?host=hh.ru',
                      'requirement': 'От 1 года коммерческой <highlighttext>разработки</highlighttext>. '
                                     'Умение структурно мыслить, а также разбивать проект на подзадачи. '
                                     'Опыт с другими языками программирования, например...',
                      'schedule': 'Полный день',
                      "vac_id": 2
                      })
    vac3 = Vacancy(**{'name': "test3",
                      'salary': 0.03,
                      'currency': "RUR",
                      'created_at': "test_date3",
                      'url': "test_url3",
                      'requirement': "test_requirement3",
                      'schedule': "test_schedule3",
                      })
    vac4 = Vacancy(**{'name': "test4",
                      'salary': 0.04,
                      'currency': "RUR",
                      'created_at': "test_date4",
                      'url': "test_url4",
                      'requirement': "test_requirement4",
                      'schedule': "test_schedule4",
                      "vac_id": 3
                      })

    vac_list1 = [vac1, vac2, vac3]
    vac_list_dict = [vac.to_dict for vac in vac_list1]
    excel_operator1 = ExcelOperator("HeadHunter", "https://hh.ru", vac_list1)

    pd1 = pd.DataFrame(vac_list_dict)
    print(pd1.shape)  # (3, 8)
    excel_operator1.add_vacancies([vac4, vac1])
